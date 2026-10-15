#!/usr/bin/env python3
"""Validate the single-sheet character/group/scene asset XLSX specification."""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


SCHEMAS = {
    "character": {
        "sheet": "人物资产",
        "master_id": "人物ID",
        "state_id": "状态ID",
        "name": "人物姓名",
        "anchor": "身份锚点",
        "stable_fields": (
            "人物画像", "对白风格", "做事与决策风格", "人际互动与情绪模式",
            "人物画像依据与置信度", "可视化气质锚点",
        ),
        "state_marker": "-S",
        "state_judgement": "状态判断（服装）",
        "required_columns": (
            "排序", "人物ID", "状态ID", "人物姓名", "状态调用标签", "状态名称",
            "状态判断（服装）", "人物画像", "对白风格", "做事与决策风格",
            "人际互动与情绪模式", "人物画像依据与置信度", "身份锚点",
            "可视化气质锚点", "服装状态", "剧情道具（左下独立展示）",
            "图片比例", "风格ID", "视觉媒介", "题材气质", "自定义风格补充",
            "画面风格", "完整生成提示词", "负面提示词",
            "来源章节/段落位置", "使用区间", "状态切换依据", "原文依据",
        ),
    },
    "group": {
        "sheet": "群像资产",
        "master_id": "群像ID",
        "state_id": "状态ID",
        "name": "群像名称",
        "anchor": "群像锚点",
        "stable_fields": ("群体身份与剧情功能",),
        "state_marker": "-S",
        "state_judgement": "状态判断（成员/服装）",
        "required_columns": (
            "排序", "群像ID", "状态ID", "群像名称", "状态调用标签", "状态名称",
            "状态判断（成员/服装）", "群体身份与剧情功能", "成员构成与人数依据",
            "具名成员人物ID/状态ID", "群像锚点", "成员与服装状态", "可视人数方案",
            "相对位置与辨识关系", "剧情道具（左下独立展示）", "图片比例",
            "风格ID", "视觉媒介", "题材气质", "自定义风格补充", "画面风格",
            "完整生成提示词", "负面提示词", "来源章节/段落位置", "使用区间",
            "状态切换依据", "原文依据",
        ),
    },
    "scene": {
        "sheet": "场景资产",
        "master_id": "场景ID",
        "state_id": "状态ID",
        "name": "场景名称",
        "anchor": "空间锚点",
        "stable_fields": (),
        "state_marker": "-T",
        "state_judgement": "状态判断（时间/空间状态）",
        "required_columns": (
            "排序", "场景ID", "状态ID", "场景名称", "状态调用标签", "状态名称",
            "状态判断（时间/空间状态）", "坐标基准", "空间锚点", "固定家具与关键道具",
            "材质与固定色板", "时间段", "空间状态（完整/被破坏等）", "四视图机位要求",
            "图片比例", "风格ID", "视觉媒介", "题材气质", "自定义风格补充",
            "画面风格", "完整生成提示词", "负面提示词",
            "来源章节/段落位置", "使用区间", "状态切换依据", "原文依据",
        ),
    },
}

MEDIA_OPTIONS = {
    "M01": "真人影视",
    "M02": "3D国漫",
    "M03": "2D国漫",
    "M04": "水墨插画",
    "M05": "日系动画",
    "M06": "美式漫画",
}
GENRE_OPTIONS = {
    "G01": "都市现实",
    "G02": "古装历史",
    "G03": "仙侠玄幻",
    "G04": "年代复古",
    "G05": "悬疑犯罪",
    "G06": "科幻未来",
}
CUSTOM_STYLE_RE = re.compile(r"STYLE-CUSTOM-\d+")
CHARACTER_LAYOUT_TOKENS = (
    "左上",
    "人物面部正面特写",
    "人物面部侧面特写",
    "右侧为人物正面、侧面、背面全身三视图",
    "人物自然站立",
    "人物双手空置",
    "手臂自然下垂或置于身体两侧",
    "严格遵循已锁定的项目视觉媒介与题材气质",
)
PROP_LAYOUT_TOKENS = ("画面左下方", "人物面部特写正下方", "完全分离", "无手部接触")
NO_PROP_LAYOUT_TOKEN = "没有道具时左下区域纯白留空"
FORBIDDEN_HANDHELD_TERMS = (
    "手握物品", "手拿物品", "手托物品", "手捧物品", "手提物品",
    "握着手机", "拿着手机", "拿着文件", "提着包", "背着包",
)
NO_PROP_VALUES = {"", "无", "没有", "不需要", "无需", "none", "n/a", "na"}
GROUP_LAYOUT_TOKENS = ("16:9横版", "纯白背景", "群体", "全身", "相对位置", "双手空置")
GROUP_PROP_TOKENS = ("左下", "独立展示", "完全分离")

SCENE_VIEW_TOKENS = (
    "第一格全景",
    "第二格中景平视",
    "第三格近景斜45度",
    "第四格全景俯拍",
)
SCENE_QUALITY_TOKENS = (
    "四宫格构图", "严格遵循已锁定的项目视觉媒介与题材气质",
    "细节丰富", "高质量", "无人",
)


def normalized(value) -> str:
    return "" if value is None else str(value).strip()


def rows_as_dicts(sheet):
    headers = [normalized(cell.value) for cell in sheet[1]]
    records = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value not in (None, "") for value in row):
            continue
        records.append((row_number, {headers[i]: row[i] for i in range(min(len(headers), len(row)))}))
    return headers, records


def missing_tokens(text: str, tokens: tuple[str, ...]) -> list[str]:
    return [token for token in tokens if token not in text]


def state_number(state_id: str, marker: str) -> int | None:
    match = re.fullmatch(rf".+{re.escape(marker)}(\d+)", state_id)
    return int(match.group(1)) if match else None


def validate_style_row(row_number, row, sheet_title: str, style: str, prompt: str) -> tuple[list[str], tuple[str, ...]]:
    errors: list[str] = []
    style_id = normalized(row.get("风格ID"))
    medium = normalized(row.get("视觉媒介"))
    genre = normalized(row.get("题材气质"))
    custom = normalized(row.get("自定义风格补充"))
    negative = normalized(row.get("负面提示词"))
    custom_is_empty = custom.lower() in NO_PROP_VALUES

    builtin_match = re.fullmatch(r"STYLE-(M\d{2})-(G\d{2})", style_id)
    if builtin_match:
        medium_id, genre_id = builtin_match.groups()
        expected_medium = MEDIA_OPTIONS.get(medium_id)
        expected_genre = GENRE_OPTIONS.get(genre_id)
        if expected_medium is None or expected_genre is None:
            errors.append(f"{sheet_title}第{row_number}行风格ID包含未知内置选项：{style_id}")
        else:
            if medium != expected_medium:
                errors.append(f"{sheet_title}第{row_number}行视觉媒介与风格ID不匹配，应为：{expected_medium}")
            if genre != expected_genre:
                errors.append(f"{sheet_title}第{row_number}行题材气质与风格ID不匹配，应为：{expected_genre}")
    elif CUSTOM_STYLE_RE.fullmatch(style_id):
        if not medium or not genre:
            errors.append(f"{sheet_title}第{row_number}行自定义风格必须填写视觉媒介和题材气质")
        if custom_is_empty:
            errors.append(f"{sheet_title}第{row_number}行自定义风格必须填写自定义风格补充")
    else:
        errors.append(f"{sheet_title}第{row_number}行风格ID格式错误：{style_id or '空'}")

    if not style:
        errors.append(f"{sheet_title}第{row_number}行画面风格为空")
    elif style not in prompt:
        errors.append(f"{sheet_title}第{row_number}行完整提示词未逐字包含画面风格")
    for field_name, value in (("视觉媒介", medium), ("题材气质", genre)):
        if not value:
            errors.append(f"{sheet_title}第{row_number}行{field_name}为空")
        elif value not in style:
            errors.append(f"{sheet_title}第{row_number}行画面风格未包含{field_name}：{value}")
    if not custom_is_empty and custom not in style:
        errors.append(f"{sheet_title}第{row_number}行画面风格未包含自定义风格补充")
    if not negative:
        errors.append(f"{sheet_title}第{row_number}行负面提示词为空")
    elif medium and medium in negative:
        errors.append(f"{sheet_title}第{row_number}行负面提示词反向排除了当前视觉媒介：{medium}")
    elif genre and genre in negative:
        errors.append(f"{sheet_title}第{row_number}行负面提示词反向排除了当前题材气质：{genre}")

    builtin_mentions = [name for name in MEDIA_OPTIONS.values() if name in style]
    if style_id.startswith("STYLE-M") and len(set(builtin_mentions)) != 1:
        errors.append(f"{sheet_title}第{row_number}行画面风格含有互斥视觉媒介：{'、'.join(builtin_mentions) or '未识别'}")

    signature = (style_id, medium, genre, custom)
    return errors, signature


def validate_sheet_format(sheet) -> list[str]:
    errors: list[str] = []
    if sheet.freeze_panes != "A2":
        errors.append(f"{sheet.title}必须冻结首行（A2）")
    if not sheet.auto_filter.ref and not any(table.autoFilter for table in sheet.tables.values()):
        errors.append(f"{sheet.title}必须开启筛选")
    for cell in sheet[1]:
        if cell.value is not None and not cell.font.bold:
            errors.append(f"{sheet.title}首行表头必须全部加粗")
            break
    return errors


def validate_kind(workbook, kind: str) -> list[str]:
    cfg = SCHEMAS[kind]
    errors: list[str] = []

    if workbook.sheetnames != [cfg["sheet"]]:
        return [f"{kind}资产Excel只能包含一个工作表“{cfg['sheet']}”，当前为：{'、'.join(workbook.sheetnames)}"]

    sheet = workbook[cfg["sheet"]]
    errors.extend(validate_sheet_format(sheet))
    headers, rows = rows_as_dicts(sheet)
    for column in cfg["required_columns"]:
        if column not in headers:
            errors.append(f"{cfg['sheet']}缺少列：{column}")
    if errors:
        return errors
    if not rows:
        return [f"{cfg['sheet']}没有资产记录"]

    state_ids: set[str] = set()
    state_labels: set[str] = set()
    entity_names: dict[str, str] = {}
    anchors: dict[str, str] = {}
    stable_values: dict[str, dict[str, str]] = {field: {} for field in cfg["stable_fields"]}
    positions: dict[str, list[int]] = {}
    state_sequences: dict[str, list[tuple[int, int]]] = {}
    block_order: list[str] = []
    project_style_signature: tuple[str, ...] | None = None

    previous_sort: int | None = None
    for position, (row_number, row) in enumerate(rows):
        master_id = normalized(row.get(cfg["master_id"]))
        state_id = normalized(row.get(cfg["state_id"]))
        entity_name = normalized(row.get(cfg["name"]))
        state_name = normalized(row.get("状态名称"))
        state_label = normalized(row.get("状态调用标签"))
        anchor = normalized(row.get(cfg["anchor"]))
        judgement = normalized(row.get(cfg["state_judgement"]))
        ratio = normalized(row.get("图片比例"))
        style = normalized(row.get("画面风格"))
        prompt = normalized(row.get("完整生成提示词"))
        source = normalized(row.get("来源章节/段落位置"))
        usage = normalized(row.get("使用区间"))
        evidence = normalized(row.get("原文依据"))

        sort_value = normalized(row.get("排序"))
        try:
            sort_number = int(sort_value)
            if previous_sort is not None and sort_number <= previous_sort:
                errors.append(f"{cfg['sheet']}第{row_number}行“排序”必须严格递增")
            previous_sort = sort_number
        except ValueError:
            errors.append(f"{cfg['sheet']}第{row_number}行“排序”必须为整数：{sort_value or '空'}")

        if not master_id:
            errors.append(f"{cfg['sheet']}第{row_number}行{cfg['master_id']}为空")
            continue
        if not entity_name:
            errors.append(f"{cfg['sheet']}第{row_number}行{cfg['name']}为空：{master_id}")
        if not state_id:
            errors.append(f"{cfg['sheet']}第{row_number}行状态ID为空：{master_id}")
        elif state_id in state_ids:
            errors.append(f"{cfg['sheet']}第{row_number}行状态ID重复：{state_id}")
        else:
            state_ids.add(state_id)

        expected_prefix = f"{master_id}{cfg['state_marker']}"
        if state_id and not state_id.startswith(expected_prefix):
            errors.append(f"{cfg['sheet']}第{row_number}行状态ID应以{expected_prefix}开头：{state_id}")
        number = state_number(state_id, cfg["state_marker"])
        if state_id and number is None:
            errors.append(f"{cfg['sheet']}第{row_number}行状态ID格式错误：{state_id}")
        elif number is not None:
            state_sequences.setdefault(master_id, []).append((row_number, number))

        expected_label = f"{state_id}｜{entity_name}｜{state_name}"
        if state_label != expected_label:
            errors.append(f"{cfg['sheet']}第{row_number}行状态调用标签错误，应为：{expected_label}")
        if state_label in state_labels:
            errors.append(f"{cfg['sheet']}第{row_number}行状态调用标签重复：{state_label}")
        elif state_label:
            state_labels.add(state_label)

        if master_id in entity_names and entity_names[master_id] != entity_name:
            errors.append(f"{cfg['sheet']}第{row_number}行同一{cfg['master_id']}的名称不一致：{master_id}")
        else:
            entity_names[master_id] = entity_name
        if not anchor:
            errors.append(f"{cfg['sheet']}第{row_number}行{cfg['anchor']}为空：{master_id}")
        elif master_id in anchors and anchors[master_id] != anchor:
            errors.append(f"{cfg['sheet']}第{row_number}行同一{cfg['master_id']}未逐字复用{cfg['anchor']}：{master_id}")
        else:
            anchors[master_id] = anchor
        if anchor and anchor not in prompt:
            errors.append(f"{cfg['sheet']}第{row_number}行完整提示词未包含{cfg['anchor']}：{master_id}")

        for field in cfg["stable_fields"]:
            value = normalized(row.get(field))
            if not value:
                errors.append(f"{cfg['sheet']}第{row_number}行{field}为空：{master_id}")
            elif master_id in stable_values[field] and stable_values[field][master_id] != value:
                errors.append(f"{cfg['sheet']}第{row_number}行同一{cfg['master_id']}未逐字复用{field}：{master_id}")
            else:
                stable_values[field][master_id] = value
            if field == "可视化气质锚点" and value and value not in prompt:
                errors.append(f"{cfg['sheet']}第{row_number}行完整提示词未包含可视化气质锚点：{master_id}")

        if master_id not in positions:
            block_order.append(master_id)
        positions.setdefault(master_id, []).append(position)
        if not state_name:
            errors.append(f"{cfg['sheet']}第{row_number}行状态名称为空：{state_id or master_id}")
        if not judgement:
            errors.append(f"{cfg['sheet']}第{row_number}行{cfg['state_judgement']}为空：{state_id or master_id}")
        if ratio != "16:9":
            errors.append(f"{cfg['sheet']}第{row_number}行图片比例必须为16:9：{ratio or '空'}")
        if not source or not usage or not evidence:
            errors.append(f"{cfg['sheet']}第{row_number}行缺少来源章节/段落位置、使用区间或原文依据：{state_id or master_id}")

        style_errors, style_signature = validate_style_row(row_number, row, cfg["sheet"], style, prompt)
        errors.extend(style_errors)
        if project_style_signature is None:
            project_style_signature = style_signature
        elif style_signature != project_style_signature:
            errors.append(f"{cfg['sheet']}第{row_number}行未复用同一项目风格组合")

        if kind == "character":
            errors.extend(validate_character_row(row_number, row, style, prompt))
        elif kind == "group":
            errors.extend(validate_group_row(row_number, row, style, prompt))
        else:
            errors.extend(validate_scene_row(row_number, row, style, prompt))

    for master_id, entity_positions in positions.items():
        expected = list(range(min(entity_positions), max(entity_positions) + 1))
        if entity_positions != expected:
            errors.append(f"同一{cfg['master_id']}的状态未连续相邻：{master_id}")

    if block_order != sorted(block_order):
        errors.append(f"{cfg['sheet']}必须按{cfg['master_id']}升序排列：当前顺序为{'、'.join(block_order)}")

    for master_id, sequence in state_sequences.items():
        numbers = [number for _, number in sequence]
        if numbers != sorted(numbers):
            errors.append(f"{master_id}的状态行未按状态ID升序排列")
        if numbers and numbers != list(range(1, len(numbers) + 1)):
            errors.append(f"{master_id}的状态ID必须从01开始连续编号：{numbers}")

    return errors


def validate_character_row(row_number, row, style: str, prompt: str) -> list[str]:
    errors: list[str] = []
    for field in ("对白风格", "做事与决策风格", "人际互动与情绪模式"):
        if not normalized(row.get(field)):
            errors.append(f"人物资产第{row_number}行{field}为空")
    profile_evidence = normalized(row.get("人物画像依据与置信度"))
    if not profile_evidence:
        errors.append(f"人物资产第{row_number}行人物画像依据与置信度为空")
    elif not any(level in profile_evidence for level in ("明确", "强推断", "弱推断")):
        errors.append(f"人物资产第{row_number}行人物画像依据与置信度必须标注明确、强推断或弱推断")
    identity_anchor = normalized(row.get("身份锚点"))
    temperament_anchor = normalized(row.get("可视化气质锚点"))
    clothing = normalized(row.get("服装状态"))
    ordered_parts = (identity_anchor, temperament_anchor, clothing, style)
    if all(ordered_parts):
        positions = [prompt.find(part) for part in ordered_parts]
        if all(position >= 0 for position in positions) and positions != sorted(positions):
            errors.append(f"人物资产第{row_number}行完整提示词必须按身份锚点→可视化气质锚点→服装状态→画面风格排序")
    missing = missing_tokens(prompt, CHARACTER_LAYOUT_TOKENS)
    if "人物双手空置" in missing and "人物双手不持有任何剧情道具" in prompt:
        missing.remove("人物双手空置")
    if missing:
        errors.append(f"人物资产第{row_number}行缺少人物版式要素：{'、'.join(missing)}")
    forbidden = [term for term in FORBIDDEN_HANDHELD_TERMS if term in prompt]
    if forbidden:
        errors.append(f"人物资产第{row_number}行包含手持剧情道具描述：{'、'.join(forbidden)}")

    if not clothing:
        errors.append(f"人物资产第{row_number}行服装状态为空")
    prop = normalized(row.get("剧情道具（左下独立展示）"))
    if prop.lower() in NO_PROP_VALUES:
        if NO_PROP_LAYOUT_TOKEN not in prompt:
            errors.append(f"人物资产第{row_number}行无道具时必须写明左下区域纯白留空")
    else:
        if prop not in prompt:
            errors.append(f"人物资产第{row_number}行完整提示词未包含剧情道具：{prop}")
        missing = missing_tokens(prompt, PROP_LAYOUT_TOKENS)
        if missing:
            errors.append(f"人物资产第{row_number}行道具布局规则不完整，缺少：{'、'.join(missing)}")
    return errors


def validate_group_row(row_number, row, style: str, prompt: str) -> list[str]:
    errors: list[str] = []
    master_id = normalized(row.get("群像ID"))
    if not re.fullmatch(r"GRP-\d{3,}", master_id):
        errors.append(f"群像资产第{row_number}行群像ID必须为GRP-001等格式")
    for field in ("成员构成与人数依据", "成员与服装状态", "可视人数方案", "相对位置与辨识关系"):
        if not normalized(row.get(field)):
            errors.append(f"群像资产第{row_number}行{field}为空")
    references = normalized(row.get("具名成员人物ID/状态ID"))
    if not references:
        errors.append(f"群像资产第{row_number}行须填写具名成员引用或“无具名成员”")
    elif references != "无具名成员" and not re.search(r"CHR-\d{3,}-S\d{2,}", references):
        errors.append(f"群像资产第{row_number}行具名成员须引用CHR-xxx-Sxx")
    anchor = normalized(row.get("群像锚点"))
    members = normalized(row.get("成员与服装状态"))
    parts = (anchor, members, style)
    if all(parts):
        positions = [prompt.find(part) for part in parts]
        if -1 in positions or positions != sorted(positions):
            errors.append(f"群像资产第{row_number}行提示词须依次包含群像锚点、成员与服装状态、画面风格")
    missing = missing_tokens(prompt, GROUP_LAYOUT_TOKENS)
    if missing:
        errors.append(f"群像资产第{row_number}行缺少群像版式要素：{'、'.join(missing)}")
    negative = normalized(row.get("负面提示词"))
    if "群像" in negative:
        errors.append(f"群像资产第{row_number}行负面提示词不能排除群像")
    if any(term in prompt for term in FORBIDDEN_HANDHELD_TERMS):
        errors.append(f"群像资产第{row_number}行包含手持剧情道具描述")
    prop = normalized(row.get("剧情道具（左下独立展示）"))
    if prop.lower() in NO_PROP_VALUES:
        if NO_PROP_LAYOUT_TOKEN not in prompt:
            errors.append(f"群像资产第{row_number}行无道具时左下区域须纯白留空")
    elif prop not in prompt or missing_tokens(prompt, GROUP_PROP_TOKENS):
        errors.append(f"群像资产第{row_number}行剧情道具须在左下独立展示并与人物完全分离")
    return errors


def validate_group_references(group_workbook, character_workbook) -> list[str]:
    errors: list[str] = []
    if character_workbook.sheetnames != ["人物资产"]:
        return ["具名成员校验需要单工作表“人物资产”Excel"]
    _, character_rows = rows_as_dicts(character_workbook["人物资产"])
    states = {normalized(row.get("状态ID")): row for _, row in character_rows}
    _, group_rows = rows_as_dicts(group_workbook["群像资产"])
    for row_number, row in group_rows:
        references = normalized(row.get("具名成员人物ID/状态ID"))
        if references == "无具名成员":
            continue
        prompt = normalized(row.get("完整生成提示词"))
        matches = re.findall(r"CHR-\d{3,}-S\d{2,}", references)
        residue = re.sub(r"CHR-\d{3,}-S\d{2,}|[、,，;；\s]+", "", references)
        if residue or not matches or len(matches) != len(set(matches)):
            errors.append(f"群像资产第{row_number}行具名成员引用格式或去重错误：{references}")
        for state_id in matches:
            character = states.get(state_id)
            if character is None:
                errors.append(f"群像资产第{row_number}行引用不存在的人物状态：{state_id}")
                continue
            for field in ("身份锚点", "服装状态"):
                value = normalized(character.get(field))
                if not value or value not in prompt:
                    errors.append(f"群像资产第{row_number}行未逐字复用{state_id}的{field}")
    return errors


def validate_scene_row(row_number, row, style: str, prompt: str) -> list[str]:
    errors: list[str] = []
    missing = missing_tokens(prompt, SCENE_VIEW_TOKENS + SCENE_QUALITY_TOKENS)
    if missing:
        errors.append(f"场景资产第{row_number}行缺少四宫格/质量要素：{'、'.join(missing)}")
    indexes = [prompt.find(token) for token in SCENE_VIEW_TOKENS]
    if all(index >= 0 for index in indexes) and indexes != sorted(indexes):
        errors.append(f"场景资产第{row_number}行四视图顺序错误")
    if "8K" not in prompt or "7680×4320" not in prompt:
        errors.append(f"场景资产第{row_number}行完整提示词必须包含8K和7680×4320")

    view_requirement = normalized(row.get("四视图机位要求"))
    missing = missing_tokens(view_requirement, SCENE_VIEW_TOKENS)
    if missing:
        errors.append(f"场景资产第{row_number}行四视图机位要求不完整：{'、'.join(missing)}")
    if not normalized(row.get("时间段")):
        errors.append(f"场景资产第{row_number}行时间段为空")
    if not normalized(row.get("空间状态（完整/被破坏等）")):
        errors.append(f"场景资产第{row_number}行空间状态为空")
    return errors


def main() -> int:
    parser = argparse.ArgumentParser(description="校验人物、群像或场景单工作表资产Excel")
    parser.add_argument("workbook", type=Path, help="人物、群像或场景提示词.xlsx")
    parser.add_argument("--kind", choices=("auto", "character", "group", "scene"), default="auto")
    parser.add_argument("--characters", type=Path, help="校验群像中的具名成员时传入人物资产Excel")
    args = parser.parse_args()

    if not args.workbook.is_file():
        print(f"ERROR: 文件不存在：{args.workbook}")
        return 2
    try:
        workbook = load_workbook(args.workbook, read_only=False, data_only=False)
    except Exception as exc:
        print(f"ERROR: 无法读取Excel：{exc}")
        return 2

    if args.kind == "auto":
        if workbook.sheetnames == ["人物资产"]:
            kind = "character"
        elif workbook.sheetnames == ["群像资产"]:
            kind = "group"
        elif workbook.sheetnames == ["场景资产"]:
            kind = "scene"
        else:
            print("ERROR: 无法识别资产类型；Excel必须只有一个“人物资产”“群像资产”或“场景资产”工作表")
            return 2
    else:
        kind = args.kind

    errors = validate_kind(workbook, kind)
    if args.characters:
        if kind != "group":
            errors.append("--characters只适用于群像资产")
        else:
            try:
                character_workbook = load_workbook(args.characters, read_only=False, data_only=False)
                errors.extend(validate_group_references(workbook, character_workbook))
            except Exception as exc:
                errors.append(f"无法读取人物资产Excel：{exc}")
    if errors:
        for error in errors:
            print(f"ERROR: {error}")
        print(f"FAIL: 共{len(errors)}项错误")
        return 1
    print(f"PASS: {args.workbook.name} 的16:9、单表结构、ID、状态相邻、锚点、风格与版式检查通过")
    return 0


if __name__ == "__main__":
    sys.exit(main())
