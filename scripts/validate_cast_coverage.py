#!/usr/bin/env python3
"""Validate that every per-section cast entry maps to a character or group asset."""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


SEPARATORS = re.compile(r"[、,，;；/／|｜\n]+")


def normalized(value) -> str:
    text = "" if value is None else str(value)
    return re.sub(r"\s+", "", text).strip()


def load_asset_names(path: Path, sheet_name: str, name_column: str, alias_column: str | None = None):
    workbook = load_workbook(path, read_only=True, data_only=False)
    if workbook.sheetnames != [sheet_name]:
        raise ValueError(f"{path.name}必须只有一个“{sheet_name}”工作表")
    sheet = workbook[sheet_name]
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
    if name_column not in headers:
        raise ValueError(f"{path.name}缺少列：{name_column}")
    name_index = headers.index(name_column)
    alias_index = headers.index(alias_column) if alias_column and alias_column in headers else None
    index: dict[str, str] = {}
    primary_names: set[str] = set()
    for row in sheet.iter_rows(min_row=2, values_only=True):
        name = "" if row[name_index] is None else str(row[name_index]).strip()
        if not name:
            continue
        primary_names.add(name)
        index[normalized(name)] = name
        if alias_index is not None and row[alias_index] is not None:
            for alias in SEPARATORS.split(str(row[alias_index])):
                alias = alias.strip()
                if alias and alias not in {"无", "无别称", "不适用"}:
                    index.setdefault(normalized(alias), name)
    return index, primary_names


def evidence_present(value) -> bool:
    if isinstance(value, str):
        return bool(value.strip())
    if isinstance(value, list):
        return any(str(item).strip() for item in value)
    return False


def main() -> int:
    parser = argparse.ArgumentParser(description="校验小说/分镜逐节出场人物是否全部映射到资产")
    parser.add_argument("manifest", type=Path, help="逐节出场覆盖清单JSON")
    parser.add_argument("--characters", type=Path, required=True, help="人物提示词Excel")
    parser.add_argument("--groups", type=Path, help="群像提示词Excel；清单含群像时必填")
    args = parser.parse_args()

    errors: list[str] = []
    try:
        data = json.loads(args.manifest.read_text(encoding="utf-8"))
    except Exception as exc:
        print(f"ERROR: 无法读取覆盖清单：{exc}")
        return 2

    if not isinstance(data, dict) or not isinstance(data.get("sections"), list):
        print('ERROR: 覆盖清单必须是包含"sections"数组的JSON对象')
        return 2

    try:
        character_index, character_names = load_asset_names(
            args.characters, "人物资产", "人物姓名", "别称/称谓"
        )
    except Exception as exc:
        print(f"ERROR: 无法读取人物资产：{exc}")
        return 2

    group_index: dict[str, str] = {}
    group_names: set[str] = set()
    if args.groups:
        try:
            group_index, group_names = load_asset_names(
                args.groups, "群像资产", "群像名称"
            )
        except Exception as exc:
            print(f"ERROR: 无法读取群像资产：{exc}")
            return 2

    seen_sections: set[str] = set()
    seen_entries: set[tuple[str, str, str]] = set()
    mapped_characters: set[str] = set()
    mapped_groups: set[str] = set()
    entry_count = 0

    for section_number, section in enumerate(data["sections"], start=1):
        if not isinstance(section, dict):
            errors.append(f"第{section_number}个section不是对象")
            continue
        section_name = str(section.get("section", "")).strip()
        source = str(section.get("source", "")).strip()
        entries = section.get("entries")
        empty_reason = str(section.get("empty_reason", "")).strip()
        label = f"{source}｜{section_name}" if source or section_name else f"section#{section_number}"
        if not section_name or not source:
            errors.append(f"{label}缺少section或source")
        section_key = normalized(label)
        if section_key in seen_sections:
            errors.append(f"小节重复：{label}")
        else:
            seen_sections.add(section_key)
        if not isinstance(entries, list):
            errors.append(f"{label}的entries必须为数组")
            continue
        if not entries and not empty_reason:
            errors.append(f"{label}没有出场条目，也未填写empty_reason")

        for item_number, item in enumerate(entries, start=1):
            entry_count += 1
            if not isinstance(item, dict):
                errors.append(f"{label}第{item_number}个条目不是对象")
                continue
            source_name = str(item.get("source_name", "")).strip()
            asset_name = str(item.get("asset_name", "")).strip()
            asset_type = str(item.get("asset_type", "")).strip().lower()
            if not source_name or not asset_name:
                errors.append(f"{label}第{item_number}个条目缺少source_name或asset_name")
                continue
            if asset_type not in {"character", "group"}:
                errors.append(f"{label}｜{source_name}的asset_type必须为character或group")
                continue
            if not evidence_present(item.get("evidence")):
                errors.append(f"{label}｜{source_name}缺少出场证据")
            entry_key = (section_key, asset_type, normalized(source_name))
            if entry_key in seen_entries:
                errors.append(f"{label}重复记录同一出场称谓：{source_name}")
            else:
                seen_entries.add(entry_key)

            if asset_type == "character":
                matched = character_index.get(normalized(asset_name))
                if matched is None:
                    errors.append(
                        f"{label}｜{source_name}未映射到人物资产：{asset_name}"
                    )
                else:
                    mapped_characters.add(matched)
            else:
                if not args.groups:
                    errors.append(
                        f"{label}｜{source_name}需要群像资产，但未提供--groups"
                    )
                    continue
                matched = group_index.get(normalized(asset_name))
                if matched is None:
                    errors.append(
                        f"{label}｜{source_name}未映射到群像资产：{asset_name}"
                    )
                else:
                    mapped_groups.add(matched)

    if not data["sections"]:
        errors.append("覆盖清单没有任何小节")
    if errors:
        for error in errors:
            print(f"ERROR: {error}")
        print(f"FAIL: 共{len(errors)}项错误")
        return 1

    print(
        "PASS: "
        f"{len(seen_sections)}节、{entry_count}个出场条目全部覆盖；"
        f"映射人物{len(mapped_characters)}/{len(character_names)}个，"
        f"群像{len(mapped_groups)}/{len(group_names)}个"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
